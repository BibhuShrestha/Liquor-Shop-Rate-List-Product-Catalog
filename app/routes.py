import csv
import io
import json
import os
import re
import uuid
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import urlretrieve

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from flask import (
    Blueprint,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    send_from_directory,
    session,
    url_for,
)
from flask_login import current_user, login_required, login_user, logout_user
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from werkzeug.utils import secure_filename

from app import db
from app.models import Category, ImportBatch, ImportRow, Product, ProductVariant, User
from config import ALLOWED_EXTENSIONS, ALLOWED_IMAGE_EXTENSIONS, BASE_DIR, UPLOAD_FOLDER

main_bp = Blueprint("main", __name__)


@main_bp.route("/")
def public_rate_list():
    query = request.args.get("q", "").strip()
    category_slug = request.args.get("category", "all")

    products_query = Product.query.join(Category).filter(Product.status == "active")
    if category_slug and category_slug != "all":
        products_query = products_query.filter(Category.slug == category_slug)
    if query:
        terms = f"%{query.lower()}%"
        products_query = products_query.filter(
            db.or_(
                db.func.lower(Product.product_name).like(terms),
                db.func.lower(Product.brand).like(terms),
                db.func.lower(Category.name).like(terms),
            )
        )

    products = products_query.order_by(Product.product_name.asc()).all()
    categories = Category.query.order_by(Category.name.asc()).all()
    return render_template("public/rate_list.html", products=products, categories=categories, query=query, active_category=category_slug)


@main_bp.route("/admin/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("main.dashboard"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user)
            return redirect(url_for("main.dashboard"))
        flash("Invalid username or password.", "danger")

    return render_template("admin/login.html")


@main_bp.route("/admin/logout")
@login_required
def logout():
    logout_user()
    flash("You have been logged out.", "success")
    return redirect(url_for("main.login"))


@main_bp.route("/admin")
@login_required
def dashboard():
    total_products = Product.query.count()
    total_categories = Category.query.count()
    total_variants = ProductVariant.query.count()
    available_products = ProductVariant.query.filter(ProductVariant.stock_status == "Available").count()
    out_of_stock = ProductVariant.query.filter(ProductVariant.stock_status == "Out of Stock").count()

    return render_template(
        "admin/dashboard.html",
        total_products=total_products,
        total_categories=total_categories,
        total_variants=total_variants,
        available_products=available_products,
        out_of_stock=out_of_stock,
    )


@main_bp.route("/admin/categories", methods=["GET", "POST"])
@login_required
def categories():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        if not name:
            flash("Category name is required.", "danger")
            return redirect(url_for("main.categories"))
        slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
        if Category.query.filter((Category.name == name) | (Category.slug == slug)).first():
            flash("Category already exists.", "warning")
            return redirect(url_for("main.categories"))
        category = Category(name=name, slug=slug)
        db.session.add(category)
        db.session.commit()
        flash("Category added successfully.", "success")
        return redirect(url_for("main.categories"))

    all_categories = Category.query.order_by(Category.name.asc()).all()
    return render_template("admin/categories.html", categories=all_categories)


@main_bp.route("/admin/products")
@login_required
def products():
    products = Product.query.order_by(Product.product_name.asc()).all()
    return render_template("admin/products.html", products=products)


@main_bp.route("/admin/products/new", methods=["GET", "POST"])
@login_required
def add_product():
    if request.method == "POST":
        product_name = request.form.get("product_name", "").strip()
        brand = request.form.get("brand", "").strip()
        category_id = request.form.get("category_id")
        description = request.form.get("description", "").strip()
        status = request.form.get("status", "active")
        category = Category.query.get(category_id)

        if not product_name or not brand or not category:
            flash("Product name, brand, and category are required.", "danger")
            return redirect(url_for("main.add_product"))

        uploaded_file = request.files.get("image")
        image_name = "placeholder.svg"
        if uploaded_file and uploaded_file.filename:
            image_name = save_uploaded_image(uploaded_file)

        product = Product(product_name=product_name, brand=brand, category_id=category.id, description=description, status=status, image=image_name)
        db.session.add(product)
        db.session.flush()

        sizes = request.form.getlist("size")
        prices = request.form.getlist("price")
        stock_statuses = request.form.getlist("stock_status")
        display_orders = request.form.getlist("display_order")

        for index, size in enumerate(sizes):
            size_value = (size or "").strip()
            if not size_value:
                continue
            price_value = parse_decimal(prices[index] if index < len(prices) else "0")
            variant = ProductVariant(
                product_id=product.id,
                size=size_value,
                price=price_value,
                discount_price=price_value,
                stock_status=(stock_statuses[index] if index < len(stock_statuses) else "Available"),
                display_order=int(display_orders[index]) if index < len(display_orders) and display_orders[index].isdigit() else index + 1,
            )
            db.session.add(variant)

        db.session.commit()
        flash("Product saved successfully.", "success")
        return redirect(url_for("main.products"))

    categories = Category.query.order_by(Category.name.asc()).all()
    return render_template("admin/product_form.html", product=None, categories=categories)


@main_bp.route("/admin/products/<int:product_id>/edit", methods=["GET", "POST"])
@login_required
def edit_product(product_id):
    product = Product.query.get_or_404(product_id)
    if request.method == "POST":
        product.product_name = request.form.get("product_name", "").strip()
        product.brand = request.form.get("brand", "").strip()
        category_id = request.form.get("category_id")
        product.category_id = int(category_id) if category_id else product.category_id
        product.description = request.form.get("description", "").strip()
        product.status = request.form.get("status", "active")

        uploaded_file = request.files.get("image")
        if uploaded_file and uploaded_file.filename:
            product.image = save_uploaded_image(uploaded_file)

        ProductVariant.query.filter_by(product_id=product.id).delete()

        sizes = request.form.getlist("size")
        prices = request.form.getlist("price")
        stock_statuses = request.form.getlist("stock_status")
        display_orders = request.form.getlist("display_order")

        for index, size in enumerate(sizes):
            size_value = (size or "").strip()
            if not size_value:
                continue
            price_value = parse_decimal(prices[index] if index < len(prices) else "0")
            db.session.add(
                ProductVariant(
                    product_id=product.id,
                    size=size_value,
                    price=price_value,
                    discount_price=price_value,
                    stock_status=(stock_statuses[index] if index < len(stock_statuses) else "Available"),
                    display_order=int(display_orders[index]) if index < len(display_orders) and display_orders[index].isdigit() else index + 1,
                )
            )

        db.session.commit()
        flash("Product updated successfully.", "success")
        return redirect(url_for("main.products"))

    categories = Category.query.order_by(Category.name.asc()).all()
    return render_template("admin/product_form.html", product=product, categories=categories)


@main_bp.route("/admin/products/<int:product_id>/delete", methods=["POST"])
@login_required
def delete_product(product_id):
    product = Product.query.get_or_404(product_id)
    db.session.delete(product)
    db.session.commit()
    flash("Product deleted successfully.", "success")
    return redirect(url_for("main.products"))


@main_bp.route("/admin/import", methods=["GET", "POST"])
@login_required
def import_products():
    if request.method == "POST":
        uploaded_file = request.files.get("file")
        if not uploaded_file or not uploaded_file.filename:
            flash("Please choose a file to import.", "danger")
            return redirect(url_for("main.import_products"))

        if not allowed_file(uploaded_file.filename):
            flash("Unsupported file format. Use CSV, XLS, XLSX, or JSON.", "danger")
            return redirect(url_for("main.import_products"))

        filename = secure_filename(uploaded_file.filename)
        file_path = os.path.join(UPLOAD_FOLDER, filename)
        uploaded_file.save(file_path)

        try:
            preview_rows = parse_import_file(file_path)
            for row in preview_rows:
                image_source = row.get("image") or ""
                if image_source and image_source.startswith(("http://", "https://")):
                    row["image"] = normalize_image_reference(image_source)
            session["import_preview"] = preview_rows
            session["import_file"] = filename
            flash(f"Imported {len(preview_rows)} rows for review.", "success")
            return redirect(url_for("main.import_products"))
        except Exception as exc:
            flash(f"Unable to parse file: {exc}", "danger")
            return redirect(url_for("main.import_products"))

    preview_rows = session.get("import_preview", [])
    return render_template("admin/import.html", preview_rows=preview_rows)


@main_bp.route("/admin/import/confirm", methods=["POST"])
@login_required
def confirm_import():
    preview_rows = session.get("import_preview")
    if not preview_rows:
        flash("No import preview available.", "warning")
        return redirect(url_for("main.import_products"))

    batch = ImportBatch(source_name=session.get("import_file", "uploaded-file"), format_type="csv", total_rows=len(preview_rows), imported_rows=0, errors=0)
    db.session.add(batch)
    db.session.flush()

    imported_count = 0
    skipped_rows = 0
    errors = 0

    for row in preview_rows:
        if row.get("status") != "valid":
            skipped_rows += 1
            errors += 1
            db.session.add(
                ImportRow(
                    batch_id=batch.id,
                    row_number=row.get("row_number", 0),
                    product_name=row.get("product_name", ""),
                    brand=row.get("brand", ""),
                    category=row.get("category", ""),
                    size=row.get("size", ""),
                    price=Decimal(str(row.get("price", 0))) if row.get("price") is not None else None,
                    image=row.get("image", ""),
                    status="invalid",
                    errors="; ".join(row.get("errors", [])),
                )
            )
            continue

        category_name = row.get("category", "")
        category = Category.query.filter_by(name=category_name).first()
        if not category:
            category = Category(name=category_name, slug=re.sub(r"[^a-z0-9]+", "-", category_name.lower()).strip("-"))
            db.session.add(category)
            db.session.flush()

        image_value = normalize_image_reference(row.get("image") or "")
        product = Product.query.filter_by(product_name=row.get("product_name"), brand=row.get("brand") or "", category_id=category.id).first()
        if not product:
            product = Product(product_name=row.get("product_name"), brand=row.get("brand") or "", category_id=category.id, status="active", image=image_value or "placeholder.svg")
            db.session.add(product)
            db.session.flush()
        elif not product.image and image_value:
            product.image = image_value

        existing_variant = ProductVariant.query.filter_by(product_id=product.id, size=row.get("size")).first()
        if existing_variant:
            existing_variant.price = Decimal(str(row.get("price", 0)))
            existing_variant.stock_status = row.get("stock_status", "Available")
            existing_variant.status = "active"
            imported_count += 1
        else:
            db.session.add(
                ProductVariant(
                    product_id=product.id,
                    size=row.get("size"),
                    price=Decimal(str(row.get("price", 0))),
                    discount_price=Decimal(str(row.get("price", 0))),
                    stock_status=row.get("stock_status", "Available"),
                    display_order=1,
                    status="active",
                )
            )
            imported_count += 1

        db.session.add(
            ImportRow(
                batch_id=batch.id,
                row_number=row.get("row_number", 0),
                product_name=row.get("product_name", ""),
                brand=row.get("brand", ""),
                category=row.get("category", ""),
                size=row.get("size", ""),
                price=Decimal(str(row.get("price", 0))) if row.get("price") is not None else None,
                image=row.get("image", ""),
                status="imported",
            )
        )

    batch.imported_rows = imported_count
    batch.errors = errors
    db.session.commit()
    session.pop("import_preview", None)
    session.pop("import_file", None)
    flash(f"Import completed. Created or updated {imported_count} variants.", "success")
    return redirect(url_for("main.dashboard"))


@main_bp.route("/export")
@login_required
def export_page():
    return render_template("admin/export.html")


@main_bp.route("/export/<export_type>")
@login_required
def export_data(export_type):
    if export_type not in {"csv", "excel", "json", "pdf"}:
        flash("Unsupported export format.", "danger")
        return redirect(url_for("main.export_page"))

    rows = []
    for product in Product.query.order_by(Product.product_name.asc()).all():
        for variant in product.variants:
            rows.append({
                "product_name": product.product_name,
                "brand": product.brand,
                "category": product.category.name if product.category else "",
                "size": variant.size,
                "price": str(variant.price),
                "stock_status": variant.stock_status,
                "image": normalize_image_reference(product.image or ""),
            })

    if export_type == "csv":
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=["product_name", "brand", "category", "size", "price", "stock_status", "image"])
        writer.writeheader()
        writer.writerows(rows)
        return send_file(io.BytesIO(output.getvalue().encode("utf-8")), mimetype="text/csv", as_attachment=True, download_name="rate_list.csv")

    if export_type == "json":
        return send_file(io.BytesIO(json.dumps(rows, indent=2).encode("utf-8")), mimetype="application/json", as_attachment=True, download_name="rate_list.json")

    if export_type == "excel":
        buffer = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "Rate List"
        headers = ["product_name", "brand", "category", "size", "price", "stock_status", "image"]
        ws.append(headers)
        for row in rows:
            image_value = row.get("image") or ""
            ws.append([
                row.get("product_name", ""),
                row.get("brand", ""),
                row.get("category", ""),
                row.get("size", ""),
                row.get("price", ""),
                row.get("stock_status", ""),
                image_value,
            ])

        for column_cells in ws.columns:
            max_length = 0
            column = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column].width = min(max(max_length + 2, 14), 40)

        ws.freeze_panes = "A2"

        product_col = ws["G"]
        for idx, cell in enumerate(product_col[1:], start=1):
            if idx >= ws.max_row:
                break
            image_value = cell.value or ""
            if not image_value:
                continue

            image_path = None
            if image_value.startswith(("http://", "https://")):
                try:
                    image_name = os.path.basename(urlparse(image_value).path) or f"img_{idx}.jpg"
                    temp_path = os.path.join(UPLOAD_FOLDER, image_name)
                    urlretrieve(image_value, temp_path)
                    image_path = temp_path
                except Exception:
                    pass
            else:
                # bare filename — check uploads dir then static/images
                for candidate in [
                    os.path.join(UPLOAD_FOLDER, image_value),
                    os.path.join(BASE_DIR, "static", "images", image_value),
                ]:
                    if os.path.exists(candidate):
                        image_path = candidate
                        break

            if not image_path:
                continue

            try:
                img = ExcelImage(image_path)
                img.width = 90
                img.height = 90
                ws.add_image(img, f"H{idx + 1}")
            except Exception:
                continue

        wb.save(buffer)
        buffer.seek(0)
        return send_file(buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="rate_list.xlsx")

    if export_type == "pdf":
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        story = [Paragraph("Liquor Shop Rate List", styles["Title"]), Spacer(1, 18)]
        data = [["Product", "Brand", "Category", "Size", "Price", "Stock"]]
        for row in rows:
            data.append([
                row["product_name"],
                row["brand"],
                row["category"],
                row["size"],
                f"Rs. {row['price']}",
                row["stock_status"],
            ])
        table = Table(data, colWidths=[130, 90, 90, 70, 70, 70])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(table)
        doc.build(story)
        buffer.seek(0)
        return send_file(buffer, mimetype="application/pdf", as_attachment=True, download_name="rate_list.pdf")

    return redirect(url_for("main.dashboard"))


@main_bp.route("/print-rate-list")
@login_required
def print_rate_list():
    products = Product.query.order_by(Product.product_name.asc()).all()
    return render_template("public/print_rate_list.html", products=products)


@main_bp.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(UPLOAD_FOLDER, filename)


@main_bp.route("/health")
def health_check():
    return jsonify({"status": "ok"})


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def normalize_image_reference(image_value):
    """Return a normalised image reference: bare filename, full URL, or 'placeholder.svg'."""
    value = str(image_value or "").strip()
    if not value or value == "placeholder.svg":
        return "placeholder.svg"
    if value.startswith(("http://", "https://")):
        return value
    # Normalise backslashes and strip any leading 'uploads/' prefix so we
    # store just the bare filename (served via the /uploads/<filename> route).
    value = value.replace("\\", "/")
    return value.removeprefix("uploads/")


def save_uploaded_image(file_storage):
    if not file_storage or not file_storage.filename:
        return "placeholder.svg"
    ext = file_storage.filename.rsplit(".", 1)[1].lower() if "." in file_storage.filename else ""
    if ext not in ALLOWED_IMAGE_EXTENSIONS:
        raise ValueError("Invalid image type. Use JPG, JPEG, PNG, or WEBP.")
    if file_storage.content_length and file_storage.content_length > 2 * 1024 * 1024:
        raise ValueError("Image file is too large. Max 2 MB.")

    safe_name = secure_filename(file_storage.filename)
    unique_name = f"{uuid.uuid4().hex}_{safe_name}"
    path = os.path.join(UPLOAD_FOLDER, unique_name)
    file_storage.save(path)
    return unique_name


def parse_decimal(value):
    if value is None or str(value).strip() == "":
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation:
        return Decimal("0")


def parse_import_file(file_path):
    ext = Path(file_path).suffix.lower()
    if ext == ".csv":
        df = pd.read_csv(file_path)
    elif ext in {".xlsx", ".xls"}:
        df = pd.read_excel(file_path)
    elif ext == ".json":
        with open(file_path, "r", encoding="utf-8") as infile:
            payload = json.load(infile)
        df = pd.DataFrame(payload)
    else:
        raise ValueError("Unsupported file format")

    if df.empty:
        return []

    rows = []
    for row_number, row in df.iterrows():
        record = {}
        for raw_key, raw_value in row.items():
            key_name = str(raw_key).strip().lower().replace(" ", "_")
            record[key_name] = raw_value

        product_name = sanitize_text(record.get("product_name") or record.get("product") or record.get("name"))
        brand = sanitize_text(record.get("brand") or record.get("company") or record.get("maker"))
        category = sanitize_text(record.get("category") or record.get("type") or record.get("category_name"))
        size = sanitize_text(record.get("size") or record.get("bottle_size") or record.get("volume"))
        price_value = record.get("price") or record.get("rate") or record.get("amount") or record.get("selling_price")
        image = sanitize_text(record.get("image") or record.get("photo") or record.get("product_image"))

        errors = []
        if not product_name:
            errors.append("Missing product name")
        if not category:
            errors.append("Missing category")
        if not size:
            errors.append("Missing size")
        if price_value is None or str(price_value).strip() == "":
            errors.append("Missing price")
        else:
            try:
                amount = Decimal(str(price_value).replace(",", "").replace("Rs.", "").replace("rs.", "").strip())
                if amount < 0:
                    raise ValueError()
            except (InvalidOperation, ValueError):
                errors.append(f"Invalid price: {price_value}")

        status = "valid" if not errors else "invalid"
        rows.append({
            "row_number": row_number + 2,
            "product_name": product_name,
            "brand": brand or "",
            "category": category or "",
            "size": size or "",
            "price": float(price_value) if price_value not in (None, "", "None") else None,
            "image": image or "",
            "stock_status": sanitize_text(record.get("stock_status")) or "Available",
            "errors": errors,
            "status": status,
        })

    return rows


def sanitize_text(value):
    if value is None:
        return ""
    return str(value).strip()
